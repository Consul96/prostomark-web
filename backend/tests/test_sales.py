from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
import uuid

from openpyxl import load_workbook
from pypdf import PdfReader
import pytest
from sqlalchemy import select
from app.models.product import Product
from app.models.sales import SalesDocumentTemplate, SalesDocumentFile
from app.services.sales.orders import calculate
from conftest import auth_headers


@pytest.fixture
def sales(db, client, seed):
    user = seed["users"]["company_admin"]
    h = auth_headers(user)
    for type in ["torg12", "upd", "invoice"]:
        db.add(
            SalesDocumentTemplate(
                company_id=None,
                document_type=type,
                name=type,
                version=1,
                status="active",
                fields={"renderer": "starter_v1"},
            )
        )
    products = []
    for n in range(2):
        p = Product(
            company_id=user.company_id,
            name=f"Крем русский {n}",
            article=f"A{n}",
            gtin=f"0460000000000{n}",
            created_by=user.id,
        )
        db.add(p)
        db.flush()
        products.append(str(p.id))
    db.commit()
    details = {
        "full_name": "ИП Тестовый Продавец",
        "legal_address": "Москва, улица Тестовая 1",
        "signatory": "Иванов Иван",
        "position": "ИП",
        "authority": "Регистрация ИП",
    }
    seller_data = {
        "kind": "ip",
        "name": "Продавец",
        "inn": "123456789012",
        "details": details,
        "vat_rules": [{"from_date": "2026-01-01", "rates": ["22", "10", "0", None]}],
    }
    seller = client.post("/api/v1/sales/sellers", json=seller_data, headers=h)
    assert seller.status_code == 201, seller.text
    customer_data = {
        "kind": "person",
        "name": "Покупатель",
        "phone": "+7 (900) 123-45-67",
        "details": {
            "legal_address": "Москва 2",
            "delivery_addresses": ["Москва 2", "Москва 3"],
        },
    }
    customer = client.post("/api/v1/sales/customers", json=customer_data, headers=h)
    assert customer.status_code == 201, customer.text
    payload = {
        "seller_id": seller.json()["id"],
        "customer_id": customer.json()["id"],
        "order_date": "2026-10-11",
        "basis": "Заказ покупателя",
        "items": [
            {
                "product_id": products[0],
                "quantity": "2",
                "price": "122",
                "vat_rate": "22",
            },
            {
                "product_id": products[1],
                "quantity": "1.005",
                "price": "10.01",
                "discount": "5",
                "vat_rate": None,
            },
        ],
    }
    return {
        "h": h,
        "payload": payload,
        "seller": seller.json(),
        "seller_data": seller_data,
        "customer": customer.json(),
        "customer_data": customer_data,
        "products": products,
    }


def order(client, sales, confirm=True):
    result = client.post(
        "/api/v1/sales/orders", headers=sales["h"], json=sales["payload"]
    )
    assert result.status_code == 201, result.text
    data = result.json()
    if confirm:
        result = client.post(
            f"/api/v1/sales/orders/{data['id']}/confirm", headers=sales["h"]
        )
        assert result.status_code == 200, result.text
        data = result.json()
    return data


def issue(client, sales, id, type="torg12", **kw):
    return client.post(
        f"/api/v1/sales/orders/{id}/documents",
        headers=sales["h"],
        json={
            "document_type": type,
            "document_date": "2026-10-11",
            "idempotency_key": str(uuid.uuid4()),
            **kw,
        },
    )


@pytest.mark.parametrize(
    "kind,inn", [("person", None), ("ip", "123456789012"), ("legal", "1234567890")]
)
def test_customer_types_search_reuse(client, sales, kind, inn):
    result = client.post(
        "/api/v1/sales/customers",
        headers=sales["h"],
        json={**sales["customer_data"], "kind": kind, "inn": inn},
    )
    assert result.status_code == 201
    assert result.json()["duplicates"]
    result = client.get("/api/v1/sales/customers?q=Покупатель", headers=sales["h"])
    assert result.json()["total"] == 2
    a, b = order(client, sales), order(client, sales)
    assert a["customer_id"] == b["customer_id"]
    assert (
        client.get(
            f"/api/v1/sales/customers/{a['customer_id']}/orders", headers=sales["h"]
        ).json()["total"]
        == 2
    )


@pytest.mark.parametrize(
    "q,p,d,r,inc,total,vat",
    [
        ("2", "122", "0", "22", True, "244.00", "44.00"),
        ("2", "100", "0", "22", False, "244.00", "44.00"),
        ("1.005", "10.01", "5", None, True, "9.56", "0.00"),
        ("1", "0.05", "0", "10", False, "0.06", "0.01"),
        ("1", "123", "0", "0", True, "123.00", "0.00"),
    ],
)
def test_decimal_rounding(q, p, d, r, inc, total, vat):
    a, b, _ = calculate(
        Decimal(q), Decimal(p), Decimal(d), Decimal(r) if r else None, inc
    )
    assert (str(a), str(b)) == (total, vat)


def test_order_totals_and_revision(client, sales):
    o = order(client, sales, False)
    assert o["total"] == "253.56" and o["vat_total"] == "44.00"
    assert len(o["items"]) == 2
    payload = {**sales["payload"], "revision": o["revision"], "comment": "Изменено"}
    assert (
        client.patch(
            f"/api/v1/sales/orders/{o['id']}", headers=sales["h"], json=payload
        ).status_code
        == 200
    )
    assert (
        client.patch(
            f"/api/v1/sales/orders/{o['id']}", headers=sales["h"], json=payload
        ).status_code
        == 409
    )
    assert (
        client.post(
            f"/api/v1/sales/orders/{o['id']}/confirm", headers=sales["h"]
        ).status_code
        == 200
    )
    assert (
        client.patch(
            f"/api/v1/sales/orders/{o['id']}", headers=sales["h"], json=payload
        ).status_code
        == 409
    )
    assert (
        client.post(
            f"/api/v1/sales/orders/{o['id']}/duplicate", headers=sales["h"]
        ).json()["status"]
        == "draft"
    )


def test_invalid_confirm_vat_and_roles(client, sales, seed):
    sales["payload"]["items"] = []
    o = order(client, sales, False)
    assert (
        client.post(
            f"/api/v1/sales/orders/{o['id']}/confirm", headers=sales["h"]
        ).status_code
        == 422
    )
    uh = auth_headers(seed["users"]["user"])
    assert (
        client.post(f"/api/v1/sales/orders/{o['id']}/confirm", headers=uh).status_code
        == 403
    )
    assert (
        client.post(
            "/api/v1/sales/sellers", headers=uh, json=sales["seller_data"]
        ).status_code
        == 403
    )
    sales["payload"]["items"] = [
        {
            "product_id": sales["products"][0],
            "quantity": "1",
            "price": "1",
            "vat_rate": "20",
        }
    ]
    assert (
        client.post(
            "/api/v1/sales/orders", headers=sales["h"], json=sales["payload"]
        ).status_code
        == 422
    )


@pytest.mark.parametrize("type", ["torg12", "upd", "invoice"])
def test_documents_download_snapshot_totals(client, sales, db, type):
    o = order(client, sales)
    result = issue(client, sales, o["id"], type)
    assert result.status_code == 201, result.text
    doc = result.json()
    assert doc["snapshot"]["total"] == o["total"]
    pdf = client.get(
        f"/api/v1/sales/documents/{doc['id']}/download?format=pdf", headers=sales["h"]
    )
    xlsx = client.get(
        f"/api/v1/sales/documents/{doc['id']}/download?format=xlsx", headers=sales["h"]
    )
    assert pdf.status_code == xlsx.status_code == 200
    text = "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf.content)).pages)
    assert "Крем русский" in text and "253.56" in text and "44.00" in text
    ws = load_workbook(BytesIO(xlsx.content)).active
    values = [str(c.value) for row in ws for c in row if c.value is not None]
    assert "253.56" in values and "44.00" in values
    assert not any(c.data_type == "f" for row in ws for c in row)
    updated = {**sales["seller_data"], "name": "Новый продавец"}
    assert (
        client.patch(
            f"/api/v1/sales/sellers/{sales['seller']['id']}",
            headers=sales["h"],
            json=updated,
        ).status_code
        == 200
    )
    assert (
        client.get(
            f"/api/v1/sales/documents/{doc['id']}/download?format=pdf",
            headers=sales["h"],
        ).content
        == pdf.content
    )
    assert (
        client.get(
            f"/api/v1/sales/documents/{doc['id']}/download?format=xlsx",
            headers=sales["h"],
        ).content
        == xlsx.content
    )
    files = db.scalars(
        select(SalesDocumentFile).where(
            SalesDocumentFile.document_id == uuid.UUID(doc["id"])
        )
    ).all()
    assert len(files) == 2 and all(f.size == len(f.content) for f in files)


def test_preview_and_missing_fields(client, sales):
    o = order(client, sales, False)
    preview = client.post(
        f"/api/v1/sales/orders/{o['id']}/documents/preview",
        headers=sales["h"],
        json={"document_type": "upd", "document_date": "2026-10-11"},
    )
    assert preview.status_code == 200
    assert "ЧЕРНОВИК" in PdfReader(BytesIO(preview.content)).pages[0].extract_text()
    assert issue(client, sales, o["id"]).status_code == 409
    assert (
        client.post(
            f"/api/v1/sales/orders/{o['id']}/confirm", headers=sales["h"]
        ).status_code
        == 200
    )
    details = sales["seller_data"]["details"].copy()
    details["signatory"] = ""
    assert (
        client.patch(
            f"/api/v1/sales/sellers/{sales['seller']['id']}",
            headers=sales["h"],
            json={**sales["seller_data"], "details": details},
        ).status_code
        == 200
    )
    result = issue(client, sales, o["id"])
    assert (
        result.status_code == 422
        and "Продавец: подписант" in result.json()["detail"]["missing_fields"]
    )


def test_cross_tenant(client, sales, seed, db):
    o = order(client, sales)
    doc = issue(client, sales, o["id"]).json()
    other = auth_headers(seed["user_b"])
    for path in [
        f"orders/{o['id']}",
        f"customers/{o['customer_id']}",
        f"sellers/{o['seller_id']}",
        f"documents/{doc['id']}/download",
        f"orders/{o['id']}/history",
    ]:
        assert client.get("/api/v1/sales/" + path, headers=other).status_code == 404
    assert (
        client.post(
            "/api/v1/sales/orders", headers=other, json=sales["payload"]
        ).status_code
        == 404
    )
    assert client.get("/api/v1/sales/orders", headers=other).json()["total"] == 0
    p = Product(
        company_id=seed["company_b"].id, name="Чужой", created_by=seed["user_b"].id
    )
    db.add(p)
    db.commit()
    payload = {
        **sales["payload"],
        "items": [{"product_id": str(p.id), "quantity": "1", "price": "1"}],
    }
    assert (
        client.post(
            "/api/v1/sales/orders", headers=sales["h"], json=payload
        ).status_code
        == 404
    )


def test_numbering_versions_concurrent(client, sales, db):
    o = order(client, sales)
    key = str(uuid.uuid4())
    with ThreadPoolExecutor(4) as pool:
        results = list(
            pool.map(
                lambda _: issue(client, sales, o["id"], idempotency_key=key), range(4)
            )
        )
    assert all(r.status_code == 201 for r in results), [r.text for r in results]
    first = results[0].json()
    assert len({r.json()["id"] for r in results}) == 1
    assert issue(client, sales, o["id"]).json()["id"] == first["id"]
    second = issue(client, sales, o["id"], new_version=True).json()
    assert second["number"] == 1 and second["version"] == 2
    assert issue(client, sales, o["id"], "upd").json()["number"] == 1
    o2 = order(client, sales)
    assert issue(client, sales, o2["id"]).json()["number"] == 2
    seller2 = client.post(
        "/api/v1/sales/sellers",
        headers=sales["h"],
        json={**sales["seller_data"], "name": "Второй ИП"},
    ).json()
    sales["payload"]["seller_id"] = seller2["id"]
    o3 = order(client, sales)
    assert issue(client, sales, o3["id"]).json()["number"] == 1
    assert (
        client.post(
            f"/api/v1/sales/orders/{o['id']}/reopen", headers=sales["h"]
        ).status_code
        == 200
    )
    sales["payload"]["seller_id"] = first["seller_id"]
    current = client.get(f"/api/v1/sales/orders/{o['id']}", headers=sales["h"]).json()
    assert (
        client.patch(
            f"/api/v1/sales/orders/{o['id']}",
            headers=sales["h"],
            json={**sales["payload"], "revision": current["revision"]},
        ).status_code
        == 200
    )
    assert (
        client.post(
            f"/api/v1/sales/orders/{o['id']}/confirm", headers=sales["h"]
        ).status_code
        == 200
    )
    assert issue(client, sales, o["id"]).status_code == 409
    assert issue(client, sales, o["id"], new_version=True).json()["version"] == 3


def test_concurrent_different_orders(client, sales):
    ids = [order(client, sales)["id"] for _ in range(4)]
    with ThreadPoolExecutor(4) as pool:
        results = list(pool.map(lambda id: issue(client, sales, id), ids))
    assert all(r.status_code == 201 for r in results), [r.text for r in results]
    assert {r.json()["number"] for r in results} == {1, 2, 3, 4}


def test_formula_injection_and_input(client, sales):
    sales["seller_data"]["details"]["full_name"] = '=HYPERLINK("bad")'
    assert (
        client.patch(
            f"/api/v1/sales/sellers/{sales['seller']['id']}",
            headers=sales["h"],
            json=sales["seller_data"],
        ).status_code
        == 200
    )
    o = order(client, sales)
    doc = issue(client, sales, o["id"]).json()
    blob = client.get(
        f"/api/v1/sales/documents/{doc['id']}/download?format=xlsx", headers=sales["h"]
    ).content
    ws = load_workbook(BytesIO(blob)).active
    assert all(c.data_type != "f" for row in ws for c in row)
    assert (
        client.post(
            "/api/v1/sales/customers",
            headers=sales["h"],
            json={**sales["customer_data"], "company_id": str(uuid.uuid4())},
        ).status_code
        == 422
    )


def test_multipage_pdf_and_totals(client, sales):
    sales["payload"]["items"] = [
        sales["payload"]["items"][0].copy() for _ in range(100)
    ]
    o = order(client, sales)
    doc = issue(client, sales, o["id"], "upd").json()
    pdf = client.get(
        f"/api/v1/sales/documents/{doc['id']}/download?format=pdf", headers=sales["h"]
    ).content
    pages = PdfReader(BytesIO(pdf)).pages
    assert len(pages) > 2
    assert sum(p.extract_text().count("Крем русский") for p in pages) == 100
    text = "\n".join(p.extract_text() for p in pages)
    assert "24400.00" in text and "4400.00" in text
    assert "Товар / артикул / GTIN" in pages[1].extract_text()


def test_document_parameter_change_requires_new_version(client, sales):
    o = order(client, sales)
    a = issue(client, sales, o["id"])
    assert a.status_code == 201
    b = issue(client, sales, o["id"], document_date="2026-10-12")
    assert b.status_code == 409
    c = issue(client, sales, o["id"], document_date="2026-10-12", new_version=True)
    assert c.status_code == 201 and c.json()["version"] == 2


def test_concurrent_draft_edits(client, sales):
    o = order(client, sales, False)

    def edit(n):
        return client.patch(
            f"/api/v1/sales/orders/{o['id']}",
            headers=sales["h"],
            json={**sales["payload"], "revision": o["revision"], "comment": str(n)},
        )

    with ThreadPoolExecutor(2) as pool:
        results = list(pool.map(edit, range(2)))
    assert sorted(r.status_code for r in results) == [200, 409]


def test_date_rules_and_amount_boundaries(client, sales):
    p = {**sales["payload"], "order_date": "2025-12-31"}
    assert (
        client.post("/api/v1/sales/orders", headers=sales["h"], json=p).status_code
        == 422
    )
    p = {
        **sales["payload"],
        "items": [
            {
                "product_id": sales["products"][0],
                "quantity": "999999999999.999",
                "price": "99999999999999.99",
            }
        ],
    }
    assert (
        client.post("/api/v1/sales/orders", headers=sales["h"], json=p).status_code
        == 422
    )
    bad = {
        **sales["seller_data"],
        "vat_rules": [
            {"from_date": "2026-01-01", "rates": [None]},
            {"from_date": "2026-02-01", "rates": ["22"]},
        ],
    }
    assert (
        client.post("/api/v1/sales/sellers", headers=sales["h"], json=bad).status_code
        == 422
    )


def test_archive_status_and_database_tenant_constraint(client, sales, seed, db):
    from sqlalchemy import text
    from sqlalchemy.exc import IntegrityError

    o = order(client, sales)
    assert (
        client.post(
            f"/api/v1/sales/orders/{o['id']}/complete", headers=sales["h"]
        ).status_code
        == 200
    )
    assert (
        client.post(
            f"/api/v1/sales/orders/{o['id']}/cancel", headers=sales["h"]
        ).status_code
        == 409
    )
    assert issue(client, sales, o["id"]).status_code == 201
    with pytest.raises(IntegrityError):
        db.execute(
            text("UPDATE sales_orders SET company_id = :company WHERE id = :id"),
            {"company": seed["company_b"].id, "id": uuid.UUID(o["id"])},
        )
        db.flush()
    db.rollback()
    customer = sales["customer"]
    assert (
        client.patch(
            f"/api/v1/sales/customers/{customer['id']}",
            headers=sales["h"],
            json={**sales["customer_data"], "archived": True},
        ).status_code
        == 200
    )
    assert (
        client.get("/api/v1/sales/customers", headers=sales["h"]).json()["total"] == 0
    )
    assert (
        client.get("/api/v1/sales/customers?archived=true", headers=sales["h"]).json()[
            "total"
        ]
        == 1
    )
    assert (
        client.post(
            "/api/v1/sales/orders", headers=sales["h"], json=sales["payload"]
        ).status_code
        == 422
    )


def test_phone_search_and_year_sequences(client, sales):
    assert (
        client.get("/api/v1/sales/customers?q=9001234567", headers=sales["h"]).json()[
            "total"
        ]
        == 1
    )
    o = order(client, sales)
    a = issue(client, sales, o["id"]).json()
    b = issue(
        client, sales, o["id"], document_date="2027-01-01", new_version=True
    ).json()
    assert a["number"] == b["number"] == 1 and a["year"] == 2026 and b["year"] == 2027
    again = issue(client, sales, o["id"], document_date="2027-01-01")
    assert again.status_code == 201 and again.json()["id"] == b["id"]
    back = issue(client, sales, o["id"], document_date="2026-10-11", new_version=True)
    assert back.status_code == 201 and back.json()["version"] == 2
    again_year = issue(
        client, sales, o["id"], document_date="2027-01-01", new_version=True
    )
    assert (
        again_year.status_code == 201
        and again_year.json()["number"] == 1
        and again_year.json()["version"] == 2
    )
    c = issue(client, sales, o["id"], document_date="2025-12-31", new_version=True)
    assert c.status_code == 422
