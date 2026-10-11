"""Normalized snapshots and versioned renderers. No external accounting side effects."""

from decimal import Decimal
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
import threading

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from sqlalchemy import select
from app.models.sales import SalesCustomer, SalesSeller, SalesDocumentTemplate
from app.services.sales.orders import scoped, validate_confirm, rates_for

TITLES = {
    "torg12": "Товарная накладная ТОРГ-12",
    "upd": "Универсальный передаточный документ",
    "invoice": "Товарная накладная",
}
FONT_PATH = Path(__file__).parent / "fonts" / "DejaVuSans.ttf"
_font_lock = threading.Lock()


def party(row):
    return {
        "id": str(row.id),
        "kind": row.kind,
        "name": row.name,
        "inn": row.inn or "",
        **row.details,
    }


def template_for(db, payload, user):
    q = select(SalesDocumentTemplate).where(
        SalesDocumentTemplate.document_type == payload.document_type,
        SalesDocumentTemplate.status == "active",
        (SalesDocumentTemplate.company_id == user.company_id)
        | SalesDocumentTemplate.company_id.is_(None),
    )
    if payload.template_id:
        q = q.where(SalesDocumentTemplate.id == payload.template_id)
    row = db.scalar(
        q.order_by(
            SalesDocumentTemplate.company_id.is_not(None).desc(),
            SalesDocumentTemplate.version.desc(),
        )
    )
    if not row:
        raise HTTPException(422, "Действующий шаблон не найден")
    if row.fields.get("renderer", "starter_v1") != "starter_v1":
        raise HTTPException(422, "Обработчик шаблона не установлен")
    return row


def document_data(db, order, payload, user, number, version, draft=False):
    seller = party(scoped(db, SalesSeller, order.seller_id, user))
    customer = party(scoped(db, SalesCustomer, order.customer_id, user))
    missing = []
    for label, p in [("Продавец", seller), ("Покупатель", customer)]:
        for field, title in [("name", "наименование"), ("legal_address", "адрес")]:
            if not p.get(field):
                missing.append(f"{label}: {title}")
        if p["kind"] != "person" and not p.get("inn"):
            missing.append(f"{label}: ИНН")
        if p["kind"] == "legal" and not p.get("kpp"):
            missing.append(f"{label}: КПП")
    for field, title in [
        ("full_name", "полное наименование"),
        ("signatory", "подписант"),
        ("position", "должность"),
        ("authority", "основание полномочий"),
    ]:
        if not seller.get(field):
            missing.append(f"Продавец: {title}")
    if not order.basis:
        missing.append("Основание передачи (договор / заказ)")
    if not order.items:
        missing.append("Товарные позиции")
    if (
        payload.document_type == "upd"
        and payload.upd_status == 1
        and any(i.vat_rate is None for i in order.items)
    ):
        missing.append(
            "УПД статус 1: проверьте применимость счёта-фактуры и ставки НДС"
        )
    if not draft:
        validate_confirm(db, order, user)
        allowed = rates_for(
            scoped(db, SalesSeller, order.seller_id, user), payload.document_date
        )
        if any(i.vat_rate not in allowed for i in order.items):
            missing.append("Ставка НДС не соответствует дате документа")
        if missing:
            raise HTTPException(422, {"missing_fields": missing})
    template = template_for(db, payload, user)
    lines = [
        {
            k: str(getattr(i, k))
            if isinstance(getattr(i, k), Decimal)
            else getattr(i, k)
            for k in [
                "name",
                "article",
                "gtin",
                "unit",
                "okei",
                "quantity",
                "price",
                "discount",
                "vat_rate",
                "total",
                "vat_amount",
            ]
        }
        for i in order.items
    ]
    return {
        "type": payload.document_type,
        "title": TITLES[payload.document_type],
        "number": number,
        "version": version,
        "date": payload.document_date.isoformat(),
        "draft": draft,
        "missing_fields": missing,
        "template_id": str(template.id),
        "template_version": template.version,
        "seller": seller,
        "customer": customer,
        "items": lines,
        "basis": order.basis,
        "order_number": order.number,
        "order_revision": order.revision,
        "delivery_address": order.delivery_address,
        "vat_included": order.vat_included,
        "total": str(order.total),
        "vat_total": str(order.vat_total),
        "discount_total": str(order.discount_total),
        "upd_status": payload.upd_status,
        "notice": "Стартовая печатная форма. Требуется бухгалтерская проверка перед реальным документооборотом.",
    }


def headings(data):
    s, c = data["seller"], data["customer"]
    rows = [
        f"{data['title']} № {data['number']} от {data['date']} · версия {data['version']}"
    ]
    if data["draft"]:
        rows.insert(0, "ЧЕРНОВИК — НЕ ВЫПУЩЕН")
    if data["type"] == "torg12":
        rows += [
            "Стартовая форма на основе ТОРГ-12 · ОКУД 0330212",
            f"ОКПО продавца: {s.get('okpo', '')}",
        ]
    if data["type"] == "upd":
        rows += [
            f"Статус УПД: {data['upd_status']}",
            "Счёт-фактура и передаточный документ"
            if data["upd_status"] == 1
            else "Передаточный документ (без счёта-фактуры)",
        ]
    for label, p in [
        ("Продавец / грузоотправитель", s),
        ("Покупатель / грузополучатель", c),
    ]:
        rows += [
            f"{label}: {p.get('full_name') or p['name']}",
            f"ИНН {p['inn']} / КПП {p.get('kpp', '')} · ОГРН {p.get('ogrn', '')}",
            f"Адрес: {p.get('legal_address', '')}",
            f"Р/с {p.get('account', '')} · {p.get('bank', '')} · БИК {p.get('bik', '')} · К/с {p.get('correspondent_account', '')}",
        ]
    rows += [
        f"Адрес отгрузки: {s.get('shipping_address', '')}",
        f"Адрес доставки: {data['delivery_address']}",
        f"Основание: {data['basis']}",
        f"Заказ: {data['order_number']} · Цены {'с НДС' if data['vat_included'] else 'без НДС'}",
    ]
    return rows


def table_rows(data):
    rows = [
        [
            "№",
            "Товар / артикул / GTIN",
            "Ед. / ОКЕИ",
            "Кол-во",
            "Цена",
            "Скидка %",
            "Без НДС",
            "НДС %",
            "НДС",
            "Всего",
        ]
    ]
    for n, i in enumerate(data["items"], 1):
        rows.append(
            [
                n,
                "\n".join(filter(None, [i["name"], i["article"], i["gtin"]])),
                i["unit"] + " / " + i["okei"],
                i["quantity"],
                i["price"],
                i["discount"],
                str(Decimal(i["total"]) - Decimal(i["vat_amount"])),
                "Без НДС" if i["vat_rate"] is None else i["vat_rate"],
                i["vat_amount"],
                i["total"],
            ]
        )
    rows.append(
        [
            "",
            "Итого",
            "",
            "",
            "",
            "",
            str(Decimal(data["total"]) - Decimal(data["vat_total"])),
            "",
            data["vat_total"],
            data["total"],
        ]
    )
    return rows


def footers(data):
    s = data["seller"]
    return [
        f"Всего к оплате: {data['total']} руб. · НДС: {data['vat_total']} руб. · Скидка: {data['discount_total']} руб.",
        f"Отпуск разрешил: {s.get('position', '')} / {s.get('signatory', '')} / подпись __________",
        f"Основание полномочий: {s.get('authority', '')}",
        "Главный бухгалтер __________________ / Отпуск произвёл __________________",
        "Груз получил / принял __________________ · Доверенность № __________ от __________",
        "Дата передачи __________ · Дата получения __________ · Подписи / печати __________",
        "Сведения о перевозке, упаковке и приложениях: __________________________________",
        "Для УПД: страна происхождения / регистрационный номер декларации (при применимости): __________________"
        if data["type"] == "upd"
        else "",
        data["notice"],
    ] + (
        ["Не заполнено: " + "; ".join(data["missing_fields"])]
        if data["draft"] and data["missing_fields"]
        else []
    )


def render_xlsx(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Документ"
    for text in headings(data):
        ws.append([text])
        ws.merge_cells(
            start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=10
        )
        ws.row_dimensions[ws.max_row].height = max(30, len(text) // 100 * 14 + 20)
    start = ws.max_row + 1
    for row in table_rows(data):
        ws.append(row)
        ws.row_dimensions[ws.max_row].height = max(
            48, (len(str(row[1])) // 38 + str(row[1]).count("\n") + 1) * 14
        )
    end = ws.max_row
    for text in filter(None, footers(data)):
        ws.append([text])
        ws.merge_cells(
            start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=10
        )
        ws.row_dimensions[ws.max_row].height = max(32, (len(text) // 100 + 1) * 14)
    for row in ws:
        for cell in row:
            # Force ALL text as text, including =, +, -, @ and leading whitespace.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="DejaVu Sans", size=10)
    for row in [1, start, end]:
        for cell in ws[row]:
            cell.font = Font(name="DejaVu Sans", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="254B69")
    for col, width in enumerate([5, 45, 13, 12, 14, 12, 15, 13, 14, 15], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f"C{start + 1}"
    ws.print_title_rows = f"{start}:{start}"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = ws.dimensions
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def render_pdf(data):
    with _font_lock:
        if "SalesDejaVu" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("SalesDejaVu", str(FONT_PATH)))
    style = ParagraphStyle("sales", fontName="SalesDejaVu", fontSize=8, leading=11)

    def p(text):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)

    out = BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        rightMargin=24,
        leftMargin=24,
        topMargin=22,
        bottomMargin=30,
    )
    story = []
    for text in headings(data):
        story += [p(text), Spacer(1, 4)]
    story += [Spacer(1, 8)]
    table = Table(
        [[p(v) for v in row] for row in table_rows(data)],
        colWidths=[34, 192, 65, 53, 65, 52, 72, 63, 63, 73],
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6EDF3")),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BCC9D4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F5F8FA")],
                ),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story += [table, Spacer(1, 12)]
    for text in filter(None, footers(data)):
        story += [p(text), Spacer(1, 6)]

    def page(canvas, doc):
        canvas.setFont("SalesDejaVu", 8)
        canvas.drawRightString(817, 14, f"Страница {doc.page}")
        if data["draft"]:
            canvas.drawString(24, 14, "ЧЕРНОВИК")

    doc.build(story, onFirstPage=page, onLaterPages=page)
    return out.getvalue()


def render(data, format):
    return render_xlsx(data) if format == "xlsx" else render_pdf(data)
